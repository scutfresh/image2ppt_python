from __future__ import annotations

import json
import re
import shutil
import subprocess

import json_repair

from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
import os
from pptx import Presentation
from rembg import remove

def align_image_to_ppt(
    source_path: Path, 
    output_path: Path, 
    target_w: int = 1920, 
    target_h: int = 1080, 
    bg_color: str = "#FFFFFF"
) -> Path:
    """
    将输入图片等比缩放并居中放置在 target_w x target_h 的画布上（补边填充）。
    """
    with Image.open(source_path) as img:
        # 转换为 RGB 以防原图带透明通道或为灰度图
        img = img.convert("RGB")
        orig_w, orig_h = img.size
        
        # 计算等比缩放系数，确保整张图都能放进目标画布
        scale = min(target_w / orig_w, target_h / orig_h)
        new_w = int(orig_w * scale)
        new_h = int(orig_h * scale)
        
        # 高质量缩放原图
        resized_img = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
        
        # 创建标准尺寸的纯色背景画布
        canvas = Image.new("RGB", (target_w, target_h), ImageColor.getrgb(bg_color))
        
        # 计算居中粘贴的坐标
        paste_x = (target_w - new_w) // 2
        paste_y = (target_h - new_h) // 2
        
        # 将缩放后的原图贴到画布中央
        canvas.paste(resized_img, (paste_x, paste_y))
        # 保存对齐后的标准化图片
        output_path.parent.mkdir(parents=True, exist_ok=True)
        canvas.save(output_path, format="PNG")
        
    return output_path

def save_bg_removed_image(input_path: Path, output_path: Path):
    """
    input_path:  原图所在的文件夹路径（抠图后的图片也将覆盖/保存到这里）
    output_path: 处理前的原图备份文件夹路径
    """
    # 确保传入的是 Path 对象
    input_folder = Path(input_path)
    output_folder = Path(output_path)
    
    # 自动创建处理前的备份文件夹
    output_folder.mkdir(parents=True, exist_ok=True)
    
    # 遍历输入文件夹中的图片
    for filename in os.listdir(input_folder):
        if filename.lower().endswith(('.png', '.jpg', '.jpeg')):
            
            # 1. 构造当前图片在输入和输出文件夹中的完整路径
            current_in_path = input_folder / filename
            current_out_path = output_folder / filename
            
            # 2. 读取处理前的原图数据
            with open(current_in_path, 'rb') as input_file:
                original_image = input_file.read()
            
            # 3. 【核心修改】将处理前的原图先保存到 output_path 文件夹
            with open(current_out_path, 'wb') as backup_file:
                backup_file.write(original_image)
            
            # 4. 执行抠图
            output_image = remove(original_image)
            

            if current_in_path.suffix.lower() in ['.jpg', '.jpeg']:
                final_save_path = current_in_path.with_suffix('.png')
                # 如果你想在生成 .png 后删掉目录里原来的 .jpg 原图，可以取消下面这行的注释：
                # current_in_path.unlink(missing_ok=True)
            else:
                final_save_path = current_in_path
                
            with open(final_save_path, 'wb') as output_file:
                output_file.write(output_image)
                
            print(f"成功！原图已备份至 {current_out_path}，抠图已保存至 {final_save_path}")

    return 0

def run_cmd(args: list[str], cwd: Path | None = None) -> str:
    result = subprocess.run(
        args,
        cwd=str(cwd) if cwd else None,
        check=True,
        capture_output=True,
        text=True,
    )
    return result.stdout.strip()


def slugify(value: str) -> str:
    text = value.strip().lower()
    text = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text or "image2ppt"


def unique_project_dir(projects_root: Path, date_prefix: str, slug: str) -> Path:
    base = projects_root / f"{date_prefix}_{slug}"
    if not base.exists():
        return base
    index = 2
    while True:
        candidate = projects_root / f"{date_prefix}_{slug}_{index}"
        if not candidate.exists():
            return candidate
        index += 1


def copy_sources(project_dir: Path, sources: list[str]) -> list[str]:
    copied: list[str] = []
    inputs_dir = project_dir / "original_inputs"
    for item in sources:
        source = Path(item).expanduser().resolve()
        if not source.exists():
            raise FileNotFoundError(f"source not found: {source}")
        target = inputs_dir / source.name
        if target.exists():
            stem = target.stem
            suffix = target.suffix
            index = 2
            while True:
                candidate = inputs_dir / f"{stem}_{index}{suffix}"
                if not candidate.exists():
                    target = candidate
                    break
                index += 1
        shutil.copy2(source, target)
        copied.append(str(target))
    return copied


def copy_prompt_file(project_dir: Path) -> str | None:
    prompt_path = Path(__file__).resolve().parent / "prompts.py"
    if not prompt_path.exists():
        return None
    inputs_dir = project_dir / "original_inputs"
    target = inputs_dir / prompt_path.name
    if target.exists():
        stem = target.stem
        suffix = target.suffix
        index = 2
        while True:
            candidate = inputs_dir / f"{stem}_{index}{suffix}"
            if not candidate.exists():
                target = candidate
                break
            index += 1
    shutil.copy2(prompt_path, target)
    return str(target)


def init_project(
    projects_root: Path,
    slug: str,
    sources: list[str],
    date: str | None,
) -> dict[str, Any]:
    projects_root.mkdir(parents=True, exist_ok=True)
    date_prefix = date or datetime.now().strftime("%Y%m%d")
    project_dir = unique_project_dir(projects_root, date_prefix, slugify(slug))
    subdirs = [
        "original_inputs",
        "component_images",
        "diagnostics",
    ]
    project_dir.mkdir(parents=True, exist_ok=True)
    for item in subdirs:
        (project_dir / item).mkdir(parents=True, exist_ok=True)

    copied_sources = copy_sources(project_dir, sources)
    copy_prompt_file(project_dir)
    return {
        "project_dir": str(project_dir),
        "original_inputs": str(project_dir / "original_inputs"),
        "component_images": str(project_dir / "component_images"),
        "diagnostics": str(project_dir / "diagnostics"),
        "sources": copied_sources,
        "manifest": str(project_dir / "manifest.json"),
        "output": str(project_dir / "output.pptx"),
        "summary": str(project_dir / "summary.json"),
    }


def run_html_svg_to_manifest(skill_root: Path, source: str, manifest_path: Path) -> None:
    script = skill_root / "scripts" / "html_svg_to_manifest.py"
    run_cmd(["python", str(script), source, "--output", str(manifest_path)])


def build_pptx(skill_root: Path, manifest: Path, output: Path, summary: Path) -> None:
    script = skill_root / "scripts" / "image2pptx.py"
    run_cmd(
        [
            "python",
            str(script),
            "build",
            "--manifest",
            str(manifest),
            "--output",
            str(output),
            "--summary",
            str(summary),
        ]
    )


def verify_pptx(path: Path) -> dict[str, int]:
    presentation = Presentation(str(path))
    slide_count = len(presentation.slides)
    shape_count = sum(len(slide.shapes) for slide in presentation.slides)
    return {"slides": slide_count, "shapes": shape_count}


def write_process_notes(path: Path, content: str) -> None:
    path.write_text(content.strip() + "\n", encoding="utf-8")


def summarize_types(items: list[dict[str, Any]], key: str) -> dict[str, int]:
    counter: Counter[str] = Counter()
    for item in items:
        value = str(item.get(key) or "unknown")
        counter[value] += 1
    return dict(counter)


def group_names_by_type(
    items: list[dict[str, Any]],
    type_key: str = "type",
    name_key: str = "name",
) -> dict[str, list[str]]:
    grouped: dict[str, list[str]] = {}
    for item in items:
        item_type = str(item.get(type_key) or "unknown")
        name = str(item.get(name_key) or "").strip()
        if not name:
            continue
        grouped.setdefault(item_type, []).append(name)
    return grouped


def build_vision_report(
    analysis: dict[str, Any],
    component_plan: dict[str, Any],
    vision_model: str,
) -> dict[str, Any]:
    titles = analysis.get("titles") or []
    body_text = analysis.get("body_text") or []
    objects = analysis.get("objects") or []
    shapes = analysis.get("shapes") or []
    background = analysis.get("background") or {}
    assets = component_plan.get("assets") or []

    needs_image_true = sum(1 for item in objects if item.get("needs_image") is True)
    needs_image_false = sum(1 for item in objects if item.get("needs_image") is False)
    transparent_true = sum(1 for item in assets if item.get("transparent") is True)
    transparent_false = sum(1 for item in assets if item.get("transparent") is False)

    return {
        "analysis": {
            "vision_model": vision_model,
            "canvas": {
                "width": analysis.get("canvas_width"),
                "height": analysis.get("canvas_height"),
            },
            "background": {
                "type": background.get("type"),
                "bbox": background.get("bbox"),
            },
            "titles": {
                "count": len(titles),
            },
            "body_text": {
                "count": len(body_text),
            },
            "objects": {
                "count": len(objects),
                "types": summarize_types(objects, "type"),
                "needs_image": {
                    "true": needs_image_true,
                    "false": needs_image_false,
                },
                "names_by_type": group_names_by_type(objects),
            },
            "shapes": {
                "count": len(shapes),
                "types": summarize_types(shapes, "type"),
            },
        },
        "component_plan": {
            "assets": {
                "count": len(assets),
                "types": summarize_types(assets, "type"),
                "transparent": {
                    "true": transparent_true,
                    "false": transparent_false,
                },
                "names_by_type": group_names_by_type(assets),
            }
        },
    }


def extract_json(text: str) -> str:
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if not match:
        raise ValueError("No JSON object found in LLM output.")
    return match.group(0)


def repair_unescaped_quotes(text: str) -> str:
    repaired: list[str] = []
    in_string = False
    escaped = False
    index = 0
    length = len(text)
    while index < length:
        char = text[index]
        if in_string:
            if escaped:
                repaired.append(char)
                escaped = False
            elif char == "\\":
                repaired.append(char)
                escaped = True
            elif char == "\"":
                lookahead = index + 1
                while lookahead < length and text[lookahead].isspace():
                    lookahead += 1
                next_char = text[lookahead] if lookahead < length else ""
                if next_char in {",", "}", "]", ":", ""}:
                    in_string = False
                    repaired.append(char)
                else:
                    repaired.append("\\\"")
            else:
                repaired.append(char)
        else:
            if char == "\"":
                in_string = True
            repaired.append(char)
        index += 1
    return "".join(repaired)


# def parse_llm_json(text: str) -> dict[str, Any]:
#     raw_json = extract_json(text)
#     try:
#         return json.loads(raw_json)
#     except json.JSONDecodeError:
#         repaired = repair_unescaped_quotes(raw_json)
#         return json.loads(repaired)
def parse_llm_json(text: str) -> dict[str, Any]:
    """
    使用 json-repair 解析并自动修复大模型返回的破损 JSON。
    即使存在未转义的双引号、漏掉的逗号，或被 Markdown ```json 包裹，也能强行救回。
    """
    try:
        # 第一层防御：直接让 json_repair 处理原始文本（它自带剥离 Markdown 标记的能力）
        parsed_data = json_repair.loads(text)
        
        # 如果模型输出了乱码导致解析出来的不是字典（比如解析成了一个纯字符串）
        if not isinstance(parsed_data, dict):
            # 第二层防御：先用你的正则强制提取出 {...} 块，再送给 json_repair 修复
            raw_json = extract_json(text)
            parsed_data = json_repair.loads(raw_json)
            
            if not isinstance(parsed_data, dict):
                raise ValueError(f"解析结果不是字典结构，而是: {type(parsed_data)}")
                
        return parsed_data
        
    except Exception as e:
        # 如果 json-repair 都无力回天，抛出异常，让 agent_workflow.py 里的 @retry_llm_call 触发重试
        raise ValueError(f"JSON 提取与自动修复彻底失败: {e}")
    
def unnormalize_bbox(bbox: dict[str, Any], canvas_w: int, canvas_h: int) -> dict[str, int]:
    """将 0-1000 的相对坐标转换为绝对像素坐标，并附带防呆容错"""
    if not bbox:
        return {"x": 0, "y": 0, "w": 0, "h": 0}
    
    # 容错：如果大模型依然输出了大于 1000 的值，说明它产生幻觉输出了绝对坐标，直接放行
    max_val = max(bbox.get("x", 0), bbox.get("y", 0), bbox.get("w", 0), bbox.get("h", 0))
    if max_val > 1000:
        print(f"警告: 发现 bbox 中的值大于 1000，疑似已是绝对坐标，跳过反归一化处理。bbox: {bbox}")
        return {
            "x": int(bbox.get("x", 0)),
            "y": int(bbox.get("y", 0)),
            "w": int(bbox.get("w", 0)),
            "h": int(bbox.get("h", 0)),
        }

    # 正常反归一化转换
    return {
        "x": int(bbox.get("x", 0) * canvas_w / 1000.0),
        "y": int(bbox.get("y", 0) * canvas_h / 1000.0),
        "w": int(bbox.get("w", 0) * canvas_w / 1000.0),
        "h": int(bbox.get("h", 0) * canvas_h / 1000.0),
    }

def convert_analysis_bboxes(analysis: dict[str, Any]) -> dict[str, Any]:
    """递归遍历 analysis JSON，将所有 bbox 还原为绝对像素坐标"""
    canvas_w = int(analysis.get("canvas_width", 1920))
    canvas_h = int(analysis.get("canvas_height", 1080))

    def _convert(item):
        if isinstance(item, dict):
            if "bbox" in item and isinstance(item["bbox"], dict):
                item["bbox"] = unnormalize_bbox(item["bbox"], canvas_w, canvas_h)
            for k, v in item.items():
                _convert(v)
        elif isinstance(item, list):
            for idx in range(len(item)):
                _convert(item[idx])

    _convert(analysis)
    return analysis

def validate_and_fix_component_plan(analysis: dict[str, Any], component_plan: dict[str, Any]) -> dict[str, Any]:
    """
    校验并修复 component_plan 中 assets 的名称和坐标，以 analysis 为基准。
    同时完美兼容 bbox 表现为 dict {"x":...} 或 list [x, y, w, h] 的两种模型幻觉行为。
    """
    
    # 1. 辅助转换函数：确保将可能为 list 形式的 bbox 归一化为标准的 dict
    def normalize_bbox_to_dict(bbox_raw: Any) -> dict:
        if isinstance(bbox_raw, dict):
            return bbox_raw
        elif isinstance(bbox_raw, list) and len(bbox_raw) >= 4:
            # 兼容模型直接吐出 [x, y, w, h] 的特殊情况
            return {
                "x": bbox_raw[0],
                "y": bbox_raw[1],
                "w": bbox_raw[2],
                "h": bbox_raw[3]
            }
        return {"x": 0, "y": 0, "w": 0, "h": 0}

    # 2. 从 analysis 中提取基准对象（需要生成图片的元素）
    reference_items = []
    
    # 提取背景 (在 agent_workflow 中被命名为 "Background Image")
    bg = analysis.get("background", {})
    if bg.get("type") == "image":
        reference_items.append({
            "name": "Background Image",
            "bbox": normalize_bbox_to_dict(bg.get("bbox"))
        })
        
    # 提取 objects 中需要生成图片的元素
    for obj in analysis.get("objects", []):
        if obj.get("needs_image"):
            reference_items.append({
                "name": obj.get("name"),
                "bbox": normalize_bbox_to_dict(obj.get("bbox"))
            })
            
    # 3. 核心比对函数（此时传入的必然已经都是经过标准化的 dict 了）
    def is_bbox_equal(b1: dict, b2: dict, tol: int = 2) -> bool:
        return (abs(b1.get("x", 0) - b2.get("x", 0)) <= tol and
                abs(b1.get("y", 0) - b2.get("y", 0)) <= tol and
                abs(b1.get("w", 0) - b2.get("w", 0)) <= tol and
                abs(b1.get("h", 0) - b2.get("h", 0)) <= tol)

    assets = component_plan.get("assets", [])
    
    # 4. 遍历 assets 并进行交叉校验
    for asset in assets:
        asset_name = asset.get("name")
        # 对当前的 asset 坐标也做一次健壮性规整
        asset_bbox = normalize_bbox_to_dict(asset.get("bbox"))
        # 将规整后的格式同步写回资产中，防止后续处理流程也因为 list 崩溃
        asset["bbox"] = asset_bbox
        
        # 检查是否完全匹配 (名字和坐标都对)
        exact_match = any(
            ref["name"] == asset_name and is_bbox_equal(ref["bbox"], asset_bbox) 
            for ref in reference_items
        )
        if exact_match:
            continue
            
        # 场景 A：名字匹配，但坐标不匹配 -> 使用 Analysis 的坐标覆盖
        name_match_ref = next((ref for ref in reference_items if ref["name"] == asset_name), None)
        if name_match_ref and not is_bbox_equal(name_match_ref["bbox"], asset_bbox):
            print(f"🔧 [校验修复] Asset '{asset_name}' 坐标幻觉。已修正: {asset_bbox} -> {name_match_ref['bbox']}")
            asset["bbox"] = name_match_ref["bbox"]
            continue
            
        # 场景 B：坐标匹配，但名字不匹配 -> 使用 Analysis 的名字覆盖
        bbox_match_ref = next((ref for ref in reference_items if is_bbox_equal(ref["bbox"], asset_bbox)), None)
        if bbox_match_ref and bbox_match_ref["name"] != asset_name:
            print(f"🔧 [校验修复] Asset 名字幻觉。坐标匹配，已将名字由 '{asset_name}' 修正为 '{bbox_match_ref['name']}'")
            asset["name"] = bbox_match_ref["name"]
            continue
            
    return component_plan