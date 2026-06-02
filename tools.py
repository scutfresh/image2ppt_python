from __future__ import annotations

import json
import re
import shutil
import subprocess

import json_repair

from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from pptx import Presentation


def run_cmd(args: list[str], cwd: Path | None = None) -> str:
    result = subprocess.run(
        args,
        cwd=str(cwd) if cwd else None,
        check=True,
        capture_output=True,
        text=True,
    )
    return result.stdout.strip()


def slugify(value: str) -> str:
    text = value.strip().lower()
    text = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text or "image2ppt"


def unique_project_dir(projects_root: Path, date_prefix: str, slug: str) -> Path:
    base = projects_root / f"{date_prefix}_{slug}"
    if not base.exists():
        return base
    index = 2
    while True:
        candidate = projects_root / f"{date_prefix}_{slug}_{index}"
        if not candidate.exists():
            return candidate
        index += 1


def copy_sources(project_dir: Path, sources: list[str]) -> list[str]:
    copied: list[str] = []
    inputs_dir = project_dir / "original_inputs"
    for item in sources:
        source = Path(item).expanduser().resolve()
        if not source.exists():
            raise FileNotFoundError(f"source not found: {source}")
        target = inputs_dir / source.name
        if target.exists():
            stem = target.stem
            suffix = target.suffix
            index = 2
            while True:
                candidate = inputs_dir / f"{stem}_{index}{suffix}"
                if not candidate.exists():
                    target = candidate
                    break
                index += 1
        shutil.copy2(source, target)
        copied.append(str(target))
    return copied


def copy_prompt_file(project_dir: Path) -> str | None:
    prompt_path = Path(__file__).resolve().parent / "prompts.py"
    if not prompt_path.exists():
        return None
    inputs_dir = project_dir / "original_inputs"
    target = inputs_dir / prompt_path.name
    if target.exists():
        stem = target.stem
        suffix = target.suffix
        index = 2
        while True:
            candidate = inputs_dir / f"{stem}_{index}{suffix}"
            if not candidate.exists():
                target = candidate
                break
            index += 1
    shutil.copy2(prompt_path, target)
    return str(target)


def init_project(
    projects_root: Path,
    slug: str,
    sources: list[str],
    date: str | None,
) -> dict[str, Any]:
    projects_root.mkdir(parents=True, exist_ok=True)
    date_prefix = date or datetime.now().strftime("%Y%m%d")
    project_dir = unique_project_dir(projects_root, date_prefix, slugify(slug))
    subdirs = [
        "original_inputs",
        "component_images",
        "diagnostics",
    ]
    project_dir.mkdir(parents=True, exist_ok=True)
    for item in subdirs:
        (project_dir / item).mkdir(parents=True, exist_ok=True)

    copied_sources = copy_sources(project_dir, sources)
    copy_prompt_file(project_dir)
    return {
        "project_dir": str(project_dir),
        "original_inputs": str(project_dir / "original_inputs"),
        "component_images": str(project_dir / "component_images"),
        "diagnostics": str(project_dir / "diagnostics"),
        "sources": copied_sources,
        "manifest": str(project_dir / "manifest.json"),
        "output": str(project_dir / "output.pptx"),
        "summary": str(project_dir / "summary.json"),
    }


def run_html_svg_to_manifest(skill_root: Path, source: str, manifest_path: Path) -> None:
    script = skill_root / "scripts" / "html_svg_to_manifest.py"
    run_cmd(["python", str(script), source, "--output", str(manifest_path)])


def build_pptx(skill_root: Path, manifest: Path, output: Path, summary: Path) -> None:
    script = skill_root / "scripts" / "image2pptx.py"
    run_cmd(
        [
            "python",
            str(script),
            "build",
            "--manifest",
            str(manifest),
            "--output",
            str(output),
            "--summary",
            str(summary),
        ]
    )


def verify_pptx(path: Path) -> dict[str, int]:
    presentation = Presentation(str(path))
    slide_count = len(presentation.slides)
    shape_count = sum(len(slide.shapes) for slide in presentation.slides)
    return {"slides": slide_count, "shapes": shape_count}


def write_process_notes(path: Path, content: str) -> None:
    path.write_text(content.strip() + "\n", encoding="utf-8")


def summarize_types(items: list[dict[str, Any]], key: str) -> dict[str, int]:
    counter: Counter[str] = Counter()
    for item in items:
        value = str(item.get(key) or "unknown")
        counter[value] += 1
    return dict(counter)


def group_names_by_type(
    items: list[dict[str, Any]],
    type_key: str = "type",
    name_key: str = "name",
) -> dict[str, list[str]]:
    grouped: dict[str, list[str]] = {}
    for item in items:
        item_type = str(item.get(type_key) or "unknown")
        name = str(item.get(name_key) or "").strip()
        if not name:
            continue
        grouped.setdefault(item_type, []).append(name)
    return grouped


def build_vision_report(
    analysis: dict[str, Any],
    component_plan: dict[str, Any],
    vision_model: str,
) -> dict[str, Any]:
    titles = analysis.get("titles") or []
    body_text = analysis.get("body_text") or []
    objects = analysis.get("objects") or []
    shapes = analysis.get("shapes") or []
    background = analysis.get("background") or {}
    assets = component_plan.get("assets") or []
    #print numbers in analysis and component_plan for debugging
    print(f"Analysis - Titles: {len(titles)}, Body Text: {len(body_text)}, Objects: {len(objects)}, Shapes: {len(shapes)}")
    print(f"Component Plan - Assets: {len(assets)}")

    needs_image_true = sum(1 for item in objects if item.get("needs_image") is True)
    needs_image_false = sum(1 for item in objects if item.get("needs_image") is False)
    transparent_true = sum(1 for item in assets if item.get("transparent") is True)
    transparent_false = sum(1 for item in assets if item.get("transparent") is False)


    return {
        "analysis": {
            "vision_model": vision_model,
            "canvas": {
                "width": analysis.get("canvas_width"),
                "height": analysis.get("canvas_height"),
            },
            "background": {
                "type": background.get("type"),
                "bbox": background.get("bbox"),
            },
            "titles": {
                "count": len(titles),
            },
            "body_text": {
                "count": len(body_text),
            },
            "objects": {
                "count": len(objects),
                "types": summarize_types(objects, "type"),
                "needs_image": {
                    "true": needs_image_true,
                    "false": needs_image_false,
                },
                "names_by_type": group_names_by_type(objects),
            },
            "shapes": {
                "count": len(shapes),
                "types": summarize_types(shapes, "type"),
            },
        },
        "component_plan": {
            "assets": {
                "count": len(assets),
                "types": summarize_types(assets, "type"),
                "transparent": {
                    "true": transparent_true,
                    "false": transparent_false,
                },
                "names_by_type": group_names_by_type(assets),
            }
        },
    }


def extract_json(text: str) -> str:
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if not match:
        raise ValueError("No JSON object found in LLM output.")
    return match.group(0)


def repair_unescaped_quotes(text: str) -> str:
    repaired: list[str] = []
    in_string = False
    escaped = False
    index = 0
    length = len(text)
    while index < length:
        char = text[index]
        if in_string:
            if escaped:
                repaired.append(char)
                escaped = False
            elif char == "\\":
                repaired.append(char)
                escaped = True
            elif char == "\"":
                lookahead = index + 1
                while lookahead < length and text[lookahead].isspace():
                    lookahead += 1
                next_char = text[lookahead] if lookahead < length else ""
                if next_char in {",", "}", "]", ":", ""}:
                    in_string = False
                    repaired.append(char)
                else:
                    repaired.append("\\\"")
            else:
                repaired.append(char)
        else:
            if char == "\"":
                in_string = True
            repaired.append(char)
        index += 1
    return "".join(repaired)


# def parse_llm_json(text: str) -> dict[str, Any]:
#     raw_json = extract_json(text)
#     try:
#         return json.loads(raw_json)
#     except json.JSONDecodeError:
#         repaired = repair_unescaped_quotes(raw_json)
#         return json.loads(repaired)
def parse_llm_json(text: str) -> dict[str, Any]:
    """
    使用 json-repair 解析并自动修复大模型返回的破损 JSON。
    即使存在未转义的双引号、漏掉的逗号，或被 Markdown ```json 包裹，也能强行救回。
    """
    try:
        # 第一层防御：直接让 json_repair 处理原始文本（它自带剥离 Markdown 标记的能力）
        parsed_data = json_repair.loads(text)
        
        # 如果模型输出了乱码导致解析出来的不是字典（比如解析成了一个纯字符串）
        if not isinstance(parsed_data, dict):
            # 第二层防御：先用你的正则强制提取出 {...} 块，再送给 json_repair 修复
            raw_json = extract_json(text)
            parsed_data = json_repair.loads(raw_json)
            
            if not isinstance(parsed_data, dict):
                raise ValueError(f"解析结果不是字典结构，而是: {type(parsed_data)}")
                
        return parsed_data
        
    except Exception as e:
        # 如果 json-repair 都无力回天，抛出异常，让 agent_workflow.py 里的 @retry_llm_call 触发重试
        raise ValueError(f"JSON 提取与自动修复彻底失败: {e}")